"""OCR abstraction. `extract_document_text` distinguishes text vs scanned docs.

Provider can be swapped later; rest of the app only depends on this interface.
"""
from __future__ import annotations

import io
import time
import os
import re
import shutil
from pathlib import Path
from sahlha.app.config import settings
from dataclasses import dataclass, field
from sahlha.app.rag.document import DocumentBlock, text_blocks, text_quality


@dataclass
class ExtractedDocument:
    text: str
    num_pages: int
    is_scanned: bool  # True when native text extraction yielded ~nothing (image-based)
    method: str       # e.g. "pypdf" | "docx" | "txt" | "ocr:tesseract" | "ocr:unavailable"
    blocks: list[DocumentBlock] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    quality: dict = field(default_factory=dict)

    def __post_init__(self):
        if not self.blocks:
            self.blocks = text_blocks(self.text)
        score = text_quality(self.text)
        if score < .65 or not self.text.strip():
            self.warnings.append("Extraction quality is low; review source or OCR output.")
        self.quality = {"method": self.method, "pages": self.num_pages,
                        "ocr_used": self.is_scanned, "extracted_character_count": len(self.text),
                        "warnings": self.warnings, "quality_indicator": score}



def _extract_pdf_native(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
    return "\n".join(parts), len(reader.pages)


def _extract_docx(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    parts.extend("\t".join(cell.text for cell in row.cells)
                 for table in doc.tables for row in table.rows)
    return "\n".join(parts)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text)
            if shape.has_table:
                parts.extend("\t".join(cell.text for cell in row.cells)
                             for row in shape.table.rows)
    return "\n".join(parts)



def discover_tesseract():
    found = shutil.which("tesseract")
    if found:
        return found
    candidates = [settings.tesseract_cmd]
    for root in (os.environ.get("ProgramFiles", ""), os.environ.get("ProgramFiles(x86)", "")):
        if root:
            candidates.append(str(Path(root) / "Tesseract-OCR" / "tesseract.exe"))
    local = os.environ.get("LOCALAPPDATA", "")
    if local:
        candidates.append(str(Path(local) / "Programs" / "Tesseract-OCR" / "tesseract.exe"))
    return next((p for p in candidates if p and Path(p).is_file()), None)


def discover_poppler():
    found = shutil.which("pdftoppm")
    if found:
        return str(Path(found).parent)
    candidates = [Path(settings.poppler_path)] if settings.poppler_path else []
    root = os.environ.get("ProgramFiles", "")
    if root:
        candidates.extend(Path(root).glob("poppler*/Library/bin"))
        candidates.extend(Path(root).glob("poppler*/bin"))
    # WinGet portable installs are outside Program Files. Discover them even
    # when this server inherited PATH before the package was installed.
    local = os.environ.get("LOCALAPPDATA", "")
    if local:
        packages = Path(local) / "Microsoft" / "WinGet" / "Packages"
        candidates.extend(sorted(packages.glob("oschwartz10612.Poppler_*/poppler-*/Library/bin"), reverse=True))
    return next((str(p) for p in candidates if any((p / name).is_file() for name in ("pdftoppm", "pdftoppm.exe"))), None)

def _ocr_languages_available() -> str:
    """Configured OCR languages with graceful fallback (ara may be missing)."""
    langs = [l.strip() for l in (settings.ocr_languages or "eng").replace(",", "+").split("+") if l.strip()]
    if not langs:
        return "eng"
    try:
        import pytesseract
        binary = discover_tesseract()
        if binary:
            pytesseract.pytesseract.tesseract_cmd = binary
        available = set(pytesseract.get_languages() or [])
        kept = [l for l in langs if l in available]
        if kept:
            # Preserve order, ensure eng fallback when everything else missing.
            return "+".join(kept)
        return "eng" if "eng" in available else langs[0]
    except Exception:
        return "+".join(langs)


def _ocr_confidence_stats(image) -> dict:
    """Aggregate Tesseract word-confidence stats (graceful when unavailable)."""
    try:
        import pytesseract
        data = pytesseract.image_to_data(image, lang=_ocr_languages_available(),
                                         output_type=pytesseract.Output.DICT)
        confs = [int(c) for c in (data.get("conf") or []) if str(c).strip() not in ("", "-1")]
        if not confs:
            return {}
        low = sum(1 for c in confs if c < 60) / len(confs)
        return {"mean_confidence": round(sum(confs) / len(confs), 1),
                "low_confidence_proportion": round(low, 3),
                "word_count": len(confs)}
    except Exception:
        return {}


def assess_page_quality(text: str) -> tuple[float, list[str]]:
    """Stronger per-page native-text quality: score + human-readable reasons.

    Signals: readable-char ratio, replacement/control chars, single-char
    fragmentation, absurd word lengths, lexical density, tiny disconnected
    lines, excessive whitespace fragments, header/footer-only pages.
    """
    from sahlha.app.rag.document import text_quality
    reasons: list[str] = []
    base = text_quality(text)
    if not (text or "").strip():
        return 0.0, ["empty page text"]
    chars = [c for c in text if not c.isspace()]
    if chars:
        readable = sum(c.isalnum() or c in ".,;:!?()[]{}-_'\"/\\|+=×÷∑∫√<>@#%؟،؛" for c in chars) / len(chars)
        if readable < 0.5:
            reasons.append(f"low readable-character ratio ({readable:.2f})")
        bad = sum(c == '\ufffd' or ord(c) < 32 and c not in "\n\t" for c in chars) / len(chars)
        if bad > 0.02:
            reasons.append(f"replacement/control characters ({bad:.2%})")
    words = re.findall(r"\w+", text, re.UNICODE)
    if words:
        if max(map(len, words)) > 40:
            reasons.append("absurd word length (possible extraction corruption)")
        singles = sum(len(w) == 1 for w in words) / len(words)
        if singles > 0.5:
            reasons.append(f"single-character fragmentation ({singles:.0%})")
        substantive = [w for w in words if len(w) >= 3]
        if len(words) >= 10 and len(substantive) / len(words) < 0.3:
            reasons.append("extremely low lexical density")
        if len(set(w.lower() for w in words)) / len(words) < 0.15 and len(words) >= 20:
            reasons.append("repeated header/footer vocabulary")
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if len(lines) >= 5:
        avg_words = sum(len(ln.split()) for ln in lines) / len(lines)
        if avg_words < 1.5:
            reasons.append("many tiny disconnected lines (layout/column-order symptoms)")
    if re.search(r"(?:\s{4,}|\t{2,})", text):
        reasons.append("excessive whitespace fragments (possible column-order loss)")
    stripped = text.strip()
    if len(stripped) < 120 and re.match(r"^(?:page\s*\d+|pg\.?\s*\d+|\d+\s*$)", stripped, re.I):
        reasons.append("page-number-only content")
    score = base
    if reasons:
        score = min(score, 0.6)
    if len(reasons) >= 3:
        score = min(score, 0.35)
    return round(score, 3), reasons


def _try_ocr_images(data: bytes, suffix: str) -> tuple[str, str] | tuple[str, str, dict]:
    """Best-effort OCR via tesseract. Returns (text, method[, confidence_info])."""
    try:
        import pytesseract
        from PIL import Image, ImageOps, ImageEnhance
    except ImportError:
        return "", "ocr:unavailable (pytesseract/Pillow not installed)"
    try:
        ocr_lang = _ocr_languages_available()

        def read_image(image, timeout):
            # Non-Pillow renderer adapters can still be passed to Tesseract.
            if not isinstance(image, Image.Image):
                text = pytesseract.image_to_string(image, lang=ocr_lang,
                    config="--psm 3 -c preserve_interword_spaces=1", timeout=timeout)
                return text, {}
            image = ImageOps.exif_transpose(image)
            try:
                orientation = pytesseract.image_to_osd(image, output_type=pytesseract.Output.DICT, timeout=min(5, timeout))
                if orientation.get("orientation_conf", 0) >= 5:
                    image = image.rotate(-orientation.get("rotate", 0), expand=True)
            except Exception:
                pass
            stats = _ocr_confidence_stats(image)
            image = ImageOps.autocontrast(ImageOps.grayscale(image))
            text = pytesseract.image_to_string(image, lang=ocr_lang,
                config="--psm 3 -c preserve_interword_spaces=1", timeout=timeout)
            return text, stats

        binary = discover_tesseract()
        if binary:
            pytesseract.pytesseract.tesseract_cmd = binary
        poppler = discover_poppler()
        if suffix == ".pdf":
            try:
                from pdf2image import convert_from_bytes, pdfinfo_from_bytes
            except ImportError:
                return "", "ocr:unavailable (pdf2image not installed)"
            deadline = time.monotonic() + 45
            pages = int(pdfinfo_from_bytes(data, timeout=10, poppler_path=poppler)["Pages"])
            texts, confs = [], []
            for page in range(1, pages + 1):
                remaining = deadline - time.monotonic()
                if remaining <= 0:
                    raise TimeoutError("OCR exceeded 45 seconds; try a smaller scan")
                images = convert_from_bytes(data, dpi=max(150, min(settings.ocr_dpi, 400)), first_page=page,
                                            last_page=page, timeout=remaining, poppler_path=poppler)
                try:
                    for image in images:
                        remaining = deadline - time.monotonic()
                        if remaining <= 0:
                            raise TimeoutError("OCR exceeded 45 seconds; try a smaller scan")
                        text, stats = read_image(image, remaining)
                        texts.append(text)
                        if stats:
                            confs.append({"page": page, **stats})
                finally:
                    for image in images:
                        image.close()
            info = {"pages": confs}
            if confs:
                means = [c["mean_confidence"] for c in confs if "mean_confidence" in c]
                if means:
                    info["mean_confidence"] = round(sum(means) / len(means), 1)
            return "\n".join(texts), "ocr:tesseract(pdf2image)", info
        with Image.open(io.BytesIO(data)) as image:
            text, stats = read_image(image, 45)
            return text, "ocr:tesseract", ({"pages": [{"page": 1, **stats}]} if stats else {})
    except Exception as exc:  # OCR is best-effort; never crash ingestion
        if type(exc).__name__ in {"TesseractNotFoundError", "PDFInfoNotInstalledError"}:
            return "", "ocr:unavailable"
        return "", "ocr:failed (try a smaller scan)"


def extract_document_text(file_bytes: bytes, filename: str) -> ExtractedDocument:
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == ".pdf":
        from pypdf import PdfReader, PdfWriter
        reader = PdfReader(io.BytesIO(file_bytes))
        blocks, texts, warnings, methods = [], [], [], []
        scanned = False
        ocr_conf_pages: list[dict] = []
        low_conf_pages: list[int] = []
        for number, page in enumerate(reader.pages, 1):
            try:
                try:
                    native = page.extract_text(extraction_mode="layout") or ""
                except TypeError:
                    native = page.extract_text() or ""
            except Exception:
                native = ""
            meaningful = re.sub(r"[^\w]", "", native)
            base_good = text_quality(native) >= .65 and (len(meaningful) >= settings.ocr_min_chars or
                (len(native.split()) >= 5 and native.rstrip().endswith((".", "!", "?", "؟", "。"))))
            page_score, page_reasons = assess_page_quality(native)
            good = base_good and page_score >= .6 and not page_reasons
            method, text = "pypdf", native
            page_conf: dict = {}
            if not good:
                scanned = True
                writer = PdfWriter()
                writer.add_page(page)
                stream = io.BytesIO()
                writer.write(stream)
                ocr_out = _try_ocr_images(stream.getvalue(), ".pdf")
                # Backwards-compatible with 2-tuple mocks/tests.
                if len(ocr_out) == 3:
                    recovered, method, info = ocr_out
                    for entry in (info or {}).get("pages", []):
                        page_conf = dict(entry, page=number)
                        ocr_conf_pages.append(page_conf)
                        if entry.get("mean_confidence", 100) < 60 or entry.get("low_confidence_proportion", 0) > 0.3:
                            low_conf_pages.append(number)
                else:
                    recovered, method = ocr_out
                if recovered.strip() and text_quality(recovered) >= text_quality(native):
                    text = recovered
                if not recovered.strip() or text_quality(text) < .65:
                    detail = ("; ".join(page_reasons[:2]) + "; " if page_reasons else "")
                    warnings.append(f"Low OCR quality on page {number} ({method}). {detail}".strip())
                elif page_conf and (page_conf.get("mean_confidence", 100) < 60):
                    warnings.append(f"Low OCR confidence on page {number} ({method}).")
                    low_conf_pages.append(number)
            texts.append(text.strip())
            page_blocks = text_blocks(text, number)
            blocks.extend(page_blocks)
            # Layout warnings: table-like text without table blocks, order risks.
            if re.search(r"\t|\|.{1,40}\|", text) and not any(b.type == "table" for b in page_blocks):
                warnings.append(f"Possible table/layout loss on page {number}.")
            tiny_lines = [ln for ln in text.splitlines() if ln.strip()]
            if len(tiny_lines) >= 8 and sum(len(ln.split()) for ln in tiny_lines) / len(tiny_lines) < 2:
                warnings.append(f"Extraction order may be unreliable on page {number}.")
            methods.append(method)
        method = methods[0] if len(set(methods)) == 1 else "hybrid:pypdf+tesseract"
        doc = ExtractedDocument("\n".join(texts), len(reader.pages), scanned, method,
                                 blocks=blocks, warnings=warnings)
        # Expose aggregate OCR confidence without storing every word.
        if ocr_conf_pages:
            means = [c.get("mean_confidence") for c in ocr_conf_pages if "mean_confidence" in c]
            doc.quality["ocr_confidence"] = {
                "pages": ocr_conf_pages,
                "mean_confidence": round(sum(means) / len(means), 1) if means else None,
                "low_confidence_pages": sorted(set(low_conf_pages)),
            }
            if low_conf_pages:
                doc.quality["warnings"] = doc.quality.get("warnings", []) + [
                    f"Low OCR confidence on page(s): {', '.join(map(str, sorted(set(low_conf_pages))))}."]
        return doc
    if suffix == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(io.BytesIO(file_bytes))
        blocks = []
        for element in document.element.body.iterchildren():
            if element.tag.endswith('}tbl'):
                table = Table(element, document)
                text = '\n'.join('\t'.join(c.text for c in row.cells) for row in table.rows)
                blocks.append(DocumentBlock(1, len(blocks), 'table', text))
            elif element.tag.endswith('}p'):
                paragraph = Paragraph(element, document)
                if not paragraph.text.strip() and not paragraph._p.xpath('.//m:oMath'):
                    continue
                style = paragraph.style.name.lower() if paragraph.style else ''
                monospace = any((run.font.name or '').lower() in {'consolas', 'courier new', 'courier', 'monospace'} for run in paragraph.runs)
                kind = ('heading' if 'heading' in style or style == 'title' else
                        'list' if 'list' in style or paragraph._p.xpath('./w:pPr/w:numPr') else
                        'code' if monospace or any(x in style for x in ('code', 'preformatted')) else 'paragraph')
                paragraph_text = ''.join(node.text or '' for node in paragraph._p.iter()
                    if node.tag.endswith('}t')) if paragraph._p.xpath('.//m:oMath') else paragraph.text
                if paragraph._p.xpath('.//m:oMath') and not paragraph.text.strip():
                    kind = 'formula'
                blocks.append(DocumentBlock(1, len(blocks), kind, paragraph_text,
                                             {'style': style, 'page_is_logical': True}))
        return ExtractedDocument('\n'.join(b.text for b in blocks), 1, False, 'docx', blocks)
    if suffix == ".pptx":
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(file_bytes))
        blocks = []
        for number, slide in enumerate(presentation.slides, 1):
            for shape in sorted(slide.shapes, key=lambda s: (s.top, s.left)):
                if shape.has_table:
                    text = '\n'.join('\t'.join(c.text for c in row.cells) for row in shape.table.rows)
                    blocks.append(DocumentBlock(number, len(blocks), 'table', text))
                elif shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            kind = 'heading' if shape == slide.shapes.title else 'list' if paragraph.level else 'paragraph'
                            blocks.append(DocumentBlock(number, len(blocks), kind, paragraph.text))
        return ExtractedDocument('\n'.join(b.text for b in blocks), len(presentation.slides), False, 'pptx', blocks)
    if suffix in {".txt", ".md", ""}:
        return ExtractedDocument(text=file_bytes.decode("utf-8", errors="ignore").strip(),
                                 num_pages=1, is_scanned=False, method="txt")
    if suffix in {".png", ".jpg", ".jpeg", ".tiff", ".bmp"}:
        ocr_out = _try_ocr_images(file_bytes, suffix)
        if len(ocr_out) == 3:
            ocr_text, method, info = ocr_out
        else:
            ocr_text, method = ocr_out
            info = {}
        doc = ExtractedDocument(text=ocr_text.strip(), num_pages=1, is_scanned=True, method=method)
        if info and info.get("pages"):
            doc.quality["ocr_confidence"] = info
        return doc
    # Fallback: try raw decode
    return ExtractedDocument(text=file_bytes.decode("utf-8", errors="ignore").strip(),
                             num_pages=1, is_scanned=False, method="raw")
